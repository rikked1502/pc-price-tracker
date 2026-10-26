"""python -B -m unittest test_report   (no DB; files go to %TEMP%)"""
import tempfile
import unittest
from html.parser import HTMLParser
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import report
import report_html


def part(slot, model, price, basis="median", n=5):
    return {"slot": slot, "model": model, "price": price, "p25": price and price * 0.9, "low": price, "basis": basis, "n": n}


def row(id, title, ask, unit, deal=False, flags=(), seen_n=1, url="https://www.facebook.com/marketplace/item/1/", **kw):
    extras = [part("monitor", "24in 144Hz", 5000.0)]
    wx = unit and unit + 5000.0
    return {"id": id, "url": url, "title": title, "ask": ask, "unit_value": unit, "buy_smart": unit and unit * 0.85,
            "extras_value": 5000.0, "with_extras": wx, "pct": wx and ask / wx, "deal": deal, "flags": sorted(flags),
            "notes": kw.get("notes", []), "parts": kw.get("parts", [part("cpu", "Ryzen 5 5600", 4500.0), part("psu", None, 1500.0, "estimated")]),
            "extras": extras, "seen_n": seen_n, "since": "2026-10-01", "last_seen": "2026-10-17", "missing": kw.get("missing", []),
            "specs": kw.get("specs") or {"cpu": "Ryzen 5 5600", "gpu": "RTX 3060", "ram": "16GB DDR4", "ssd": "512GB",
                                         "unknown": kw.get("unknown", [])}}


ROWS = [
    row(1, "Ryzen 5600 RTX 3060 gaming PC", 25000.0, 30000.0, deal=True),
    row(2, "i5 12400F RX 6600 set", 30000.0, 29000.0, flags={"check"}, notes=["gpu size unknown: RTX 3060 (8/12GB)"]),
    row(3, "mystery build", 40000.0, None, flags={"can't value"}, parts=[part("gpu", "GTX 1070 Ti", None, "none", 0)],
        missing=["GTX 1070 Ti: no listings yet"]),
    row(6, "<b>drum set</b>", 21000.0, None, flags={"can't value"}, specs={"cpu": None, "gpu": None, "unknown": []}),
    row(4, "Repost i7 build", 35000.0, 36000.0, flags={"estimated", "manual"}, seen_n=3),
    row(5, '<script>alert(1)</script> 🔥 BUDOL PC "fast" & cheap', 20000.0, 26000.0, deal=True,
        url="javascript:alert(1)", unknown=["<img src=x onerror=alert(2)>"]),
]


class Tags(HTMLParser):
    def __init__(self):
        super().__init__()
        self.stack, self.refs, self.scripts = [], [], 0

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        self.refs += [a[k] for k in ("href", "src") if k in a]
        self.scripts += tag == "script"
        if tag not in ("meta", "input", "br", "img"):
            self.stack.append(tag)

    def handle_endtag(self, tag):
        assert self.stack and self.stack.pop() == tag, f"unbalanced </{tag}>"


class ReportTest(unittest.TestCase):
    def test_html(self):
        out = Path(tempfile.gettempdir()) / "pc_vs_parts_test.html"
        page = Path(report_html.write(ROWS, out)).read_text(encoding="utf-8")
        p = Tags()
        p.feed(page)
        self.assertEqual(p.stack, [])
        self.assertEqual(p.scripts, 1)  # only our inline script, the listing's <script> stays text
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt; 🔥", page)
        self.assertNotIn("<img src=x", page)
        self.assertNotIn("javascript:", page)
        self.assertTrue(all(r.startswith("https://www.facebook.com/") for r in p.refs), p.refs)
        self.assertIn("₱25,000", page)
        self.assertIn("3× since 2026-10-01", page)
        self.assertIn("<strong>4</strong> valued", page)  # row 3 can't be valued: it goes to "Missing a price", not the table
        self.assertIn("Missing a price (1)", page)
        # row 1: cpu 4,500 + psu 1,500 = 6,000 median; extras 5,000; grand 11,000. Lowest falls back to the price
        self.assertIn('Parts total</th><td class="n">₱6,000</td><td class="n">₱5,400</td><td class="n">₱6,000</td>', page)
        self.assertIn('Total with extras</th><td class="n">₱11,000</td><td class="n">₱9,900</td><td class="n">₱11,000</td>', page)
        self.assertIn("GTX 1070 Ti: no listings yet", page)
        self.assertIn("Most needed: GTX 1070 Ti (1)", page)
        self.assertIn("No specs listed (1)", page)
        self.assertIn("&lt;b&gt;drum set", page)
        print("\n  sample HTML:", out)
        empty = Path(report_html.write([], out.with_name("pc_vs_parts_empty.html"))).read_text(encoding="utf-8")
        self.assertIn("No valued builds yet", empty)

    def test_excel(self):
        df = report.pc_sheet(ROWS)
        self.assertEqual(list(df.columns), report.PC_COLS)
        self.assertEqual(list(df.deal), ["DEAL", "", "", "", "", "DEAL"])
        self.assertIn("CPU Ryzen 5 5600 ₱4,500 median·5; PSU ₱1,500 est", df.parts[0])
        self.assertIn("GPU GTX 1070 Ti can't value", df.parts[2])
        self.assertEqual(df.missing[2], "GTX 1070 Ti: no listings yet")
        out = Path(tempfile.gettempdir()) / "pc_vs_parts_test.xlsx"
        with pd.ExcelWriter(out, engine="openpyxl") as xw:
            report._sheet(xw, report.PC, df)
            report._sheet(xw, "Empty", report.pc_sheet([]))
        ws = load_workbook(out)[report.PC]
        self.assertEqual(ws["A1"].fill.fgColor.rgb, "001E40AF")
        self.assertEqual(ws["B2"].number_format, "0%")
        url = ws.cell(2, report.PC_COLS.index("url") + 1)
        self.assertEqual((url.value, url.hyperlink.target), ("open", ROWS[0]["url"]))
        self.assertEqual(ws["H2"].fill.fgColor.rgb, "00DCFCE7")  # deal row tinted across
        self.assertEqual(ws.cell(3, report.PC_COLS.index("flags") + 1).fill.fgColor.rgb, "00FEF3C7")
        self.assertEqual(load_workbook(out)["Empty"].max_row, 1)
        print("\n  sample xlsx:", out)


if __name__ == "__main__":
    unittest.main()
