from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

import price
import report_html
import trends_html
import value

OUT = Path(__file__).with_name("output")
ARCHIVE = OUT / "archive"


def _archive():
    """Older reports move to output/archive/ so output/ holds only the latest run. A file open in Excel stays put."""
    ARCHIVE.mkdir(parents=True, exist_ok=True)
    for f in [*OUT.glob("masterlist_*.xlsx"), *OUT.glob("pc_vs_parts_*.html")]:
        try:
            f.replace(ARCHIVE / f.name)
        except PermissionError:
            print(f"  {f.name} is open, left in output/")
# palette: ui-ux-pro-max "Analytics Dashboard" (blue data, green = cheap end, amber = expensive end)
HEADER_FILL, HEADER_FONT = PatternFill("solid", fgColor="1E40AF"), Font(bold=True, color="FFFFFF")
LOW = (PatternFill("solid", fgColor="DCFCE7"), Font(color="166534", bold=True))
HIGH = (PatternFill("solid", fgColor="FEF3C7"), Font(color="92400E", bold=True))
TINT = {"lowest": LOW, "buy_under": LOW, "highest": HIGH, "sell_at": HIGH, "margin": LOW}
PESO = "₱#,##0"
FORMATS = {**{c: PESO for c in ["lowest", "highest", "buy_under", "median", "sell_at", "margin", "price",
                                "ask", "unit_value", "buy_smart", "with_extras", "extras_value"]},
           "margin_pct": "0%", "pct": "0%", **{c: "yyyy-mm-dd" for c in ["lowest_date", "highest_date", "last_seen", "date"]}}
WIDTHS = {"model": 26, "title": 60, "category": 12, "sources": 24, "flags": 14, "city": 22,
          "deal": 8, "pct": 8, "cpu": 22, "gpu": 22, "parts": 90, "notes": 40, "seen": 20}
LINKS = {"lowest_link", "highest_link", "url"}
PC = "PC vs Parts"
PC_COLS = ["deal", "pct", "ask", "unit_value", "buy_smart", "with_extras", "extras_value", "title",
           "cpu", "gpu", "ram", "ssd", "parts", "flags", "notes", "missing", "seen", "url"]
BASIS = {"estimated": "est", "manual": "manual", "none": "can't value"}


def _date(s):
    return s.map(lambda t: datetime.fromtimestamp(t).date())


def _part(p):
    basis = f"median·{p['n']}" if p["basis"] == "median" else BASIS.get(p["basis"], p["basis"])
    price_ = "" if p["price"] is None else f"₱{p['price']:,.0f}"
    return " ".join(x for x in (p["slot"].upper(), p["model"] or "", price_, basis) if x)


def pc_sheet(rows):
    return pd.DataFrame([{
        "deal": "DEAL" if r["deal"] else "", "pct": r["pct"], "ask": r["ask"], "unit_value": r["unit_value"],
        "buy_smart": r["buy_smart"], "with_extras": r["with_extras"], "extras_value": r["extras_value"],
        "title": r["title"], **{k: report_html.text(r["specs"].get(k)) for k in ("cpu", "gpu", "ram", "ssd")},
        "parts": "; ".join(_part(p) for p in r["parts"] + r["extras"]), "flags": ", ".join(r["flags"]),
        "notes": "; ".join(r["notes"]), "missing": "; ".join(r.get("missing") or []), "seen": f"{r['seen_n']}× since {r['since']}" if r["seen_n"] > 1 else "",
        "url": r["url"]} for r in rows], columns=PC_COLS)


def _sheet(xw, name, sheet):
    sheet.to_excel(xw, sheet_name=name, index=False)
    ws = xw.sheets[name]
    ws.freeze_panes = "B2"  # header + first column stay put while scrolling
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        header = col[0].value
        col[0].fill, col[0].font = HEADER_FILL, HEADER_FONT
        ws.column_dimensions[col[0].column_letter].width = WIDTHS.get(header, 13)
        for c in col[1:]:
            if header in LINKS and c.value:
                c.hyperlink, c.value, c.style = c.value, "open", "Hyperlink"
            if header in FORMATS:
                c.number_format = FORMATS[header]
            if header in TINT:
                c.fill, c.font = TINT[header]
    if name == PC:  # whole deal row green, any flag amber so it gets checked by hand
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "DEAL":
                for c in row:
                    c.fill = LOW[0]
            flags = row[PC_COLS.index("flags")]
            if flags.value:
                flags.fill, flags.font = HIGH


def write(con):
    obs = price.observations(con)
    ml = price.masterlist(obs)
    fl = price.flips(ml)
    for d in (ml, fl):
        for c in ["lowest_date", "highest_date", "last_seen"]:
            d[c] = _date(d[c])
    hist = obs.assign(date=_date(obs.seen_at)).sort_values("seen_at", ascending=False)
    unmatched = pd.read_sql("""select source, title, price, category, flags, last_seen, url from listings
                               where model is null and norm_source is not null order by last_seen desc""", con)
    unmatched["last_seen"] = _date(unmatched.last_seen)
    builds = value.value_builds(con)
    sheets = {
        "Masterlist": ml[["category", "model", "listings", "lowest", "lowest_date", "highest", "highest_date",
                          "recent", "buy_under", "median", "sell_at", "sources", "last_seen", "lowest_link", "highest_link"]],
        PC: pc_sheet(builds),  # second, not first: parts-only runs leave it empty and Masterlist stays the landing tab
        "Flips": fl[["category", "model", "recent", "buy_under", "sell_at", "margin", "margin_pct", "lowest", "highest", "last_seen"]],
        "History": hist[["date", "source", "category", "model", "price", "condition", "title", "city", "url"]],
        "Unmatched": unmatched,
    }
    OUT.mkdir(exist_ok=True)
    _archive()
    ts =f"{datetime.now():%Y-%m-%d_%H%M}"  # timestamped so an open Excel file never blocks the write
    path = OUT / f"masterlist_{ts}.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, sheet in sheets.items():
            _sheet(xw, name, sheet)
    print("  HTML:", report_html.write(builds, OUT / f"pc_vs_parts_{ts}.html", trends_html.section(ml, obs)))
    return path
